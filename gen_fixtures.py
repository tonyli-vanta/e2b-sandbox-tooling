#!/usr/bin/env python3
"""Generate a deterministic fixture set for the E2B capability-matrix commands.
Each file has known content so expected extraction output is predictable."""
import os
OUT = "/tmp/e2b-fixtures"
os.makedirs(OUT, exist_ok=True)

# --- CSV (pandas) ---
open(f"{OUT}/sample.csv", "w").write("id,name,amount,owner\n1,Alpha,10,a@example.com\n2,Beta,20,b@example.com\n3,Gamma,30,c@example.com\n")

# --- HTML (bs4 + lxml) ---
open(f"{OUT}/sample.html", "w").write(
    "<html><head><title>E2B Fixture</title></head><body>"
    "<h1>E2B fixture HTML</h1><p>Hello <b>world</b>. Value 42.</p>"
    "<table><tr><th>k</th><th>v</th></tr><tr><td>a</td><td>1</td></tr></table>"
    "</body></html>")

# --- XLSX (pandas + openpyxl) with two gotchas: a mailto hyperlink + phantom empty rows ---
from openpyxl import Workbook
from openpyxl.styles import PatternFill
wb = Workbook(); ws = wb.active; ws.title = "Controls"
ws.append(["ID", "Name", "Owner"])
ws.append(["C-1", "Access Control", "alice@example.com"])
ws.append(["C-2", "Encryption", "bob@example.com"])
ws["C2"].hyperlink = "mailto:alice@example.com"   # hyperlink gotcha (pandas drops this)
ws["C3"].hyperlink = "mailto:bob@example.com"
# phantom-row trap: format a far-down empty cell so max_row reports ~500 with no data
ws["A500"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
wb.save(f"{OUT}/sample.xlsx")

# --- Legacy XLS (pandas + xlrd to read; xlwt to write) ---
import xlwt
w = xlwt.Workbook(); s = w.add_sheet("Sheet1")
for r, row in enumerate([["ID", "Name", "Amount"], ["1", "Alpha", 10], ["2", "Beta", 20]]):
    for c, val in enumerate(row):
        s.write(r, c, val)
w.save(f"{OUT}/sample_legacy.xls")

# --- Text PDF (pdftotext / pypdf) — real selectable text layer ---
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
cv = canvas.Canvas(f"{OUT}/text.pdf", pagesize=letter)
for i, line in enumerate([
    "E2B fixture: text PDF.",
    "This line is real selectable text, extractable by pdftotext.",
    "Reference number 12345 and keyword COMPLIANCE.",
]):
    cv.drawString(72, 740 - i * 24, line)
cv.showPage(); cv.save()

# --- Scanned / image-only PDF (OCR gap) — text rendered as pixels, no text layer ---
from PIL import Image, ImageDraw
img = Image.new("RGB", (1100, 300), "white")
d = ImageDraw.Draw(img)
d.text((30, 130), "This text is an IMAGE (no text layer) - needs OCR. Code 67890.", fill="black")
img.save(f"{OUT}/scanned.pdf", "PDF", resolution=150.0)

# --- Image with text (PNG) — Pillow opens it; OCR gap for text ---
img2 = Image.new("RGB", (800, 200), "white")
ImageDraw.Draw(img2).text((20, 90), "Text inside a PNG image. Token ABCDE.", fill="black")
img2.save(f"{OUT}/image_text.png")

# --- DOCX (python-docx) ---
from docx import Document
doc = Document()
doc.add_heading("E2B fixture DOCX", level=0)
doc.add_paragraph("First paragraph of the Word document.")
doc.add_paragraph("Second paragraph with data value 42 and keyword AUDIT.")
doc.save(f"{OUT}/sample.docx")

# --- Legacy .doc (libreoffice converts from sample.docx; antiword fixture) ---
# python-docx cannot write the binary .doc format and there is no pure-Python alternative,
# so we shell libreoffice on the dev box. Skips cleanly when libreoffice is not installed.
import shutil, subprocess
if shutil.which("libreoffice"):
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "doc", "--outdir", OUT, f"{OUT}/sample.docx"],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
else:
    print("libreoffice not found on PATH; skipping sample.doc generation (install libreoffice to produce the antiword fixture).")

# --- PPTX (python-pptx) — gap in sandbox, fixture proves the gap ---
from pptx import Presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "E2B fixture PPTX"
slide.placeholders[1].text = "Subtitle text with keyword DECK and number 999."
prs.save(f"{OUT}/sample.pptx")

print("Generated fixtures in", OUT)
for f in sorted(os.listdir(OUT)):
    print(f"  {f:22s} {os.path.getsize(os.path.join(OUT,f)):>8d} bytes")
